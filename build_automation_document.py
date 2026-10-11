from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"C:\Users\JOSE.LAPTOP-SNQTVDID\Desktop\MUSICAL\docs\Manual_tecnico_automatizacion_Mysic.docx")

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(9.5)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.08
for name, size, before, after in [('Title', 25, 0, 12), ('Heading 1', 17, 14, 7), ('Heading 2', 13, 11, 5), ('Heading 3', 10.5, 8, 3)]:
    s = styles[name]
    s.font.name = 'Aptos Display'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0,0,0)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

if 'Code Block' not in styles:
    cs = styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
else:
    cs = styles['Code Block']
cs.font.name = 'Consolas'; cs.font.size = Pt(7.5); cs.font.color.rgb = RGBColor(32, 41, 51)
cs.paragraph_format.left_indent = Inches(0.2); cs.paragraph_format.right_indent = Inches(0.2)
cs.paragraph_format.space_before = Pt(4); cs.paragraph_format.space_after = Pt(7); cs.paragraph_format.line_spacing = 1.0

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color='D9D9D9'):
    tblPr = table._tbl.tblPr
    el = tblPr.find(qn('w:tblBorders'))
    if el is None: el = OxmlElement('w:tblBorders'); tblPr.append(el)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement(f'w:{edge}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),color); el.append(x)

def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); m = tcPr.first_child_found_in('w:tcMar')
    if m is None: m=OxmlElement('w:tcMar'); tcPr.append(m)
    for tag,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        e=m.find(qn(f'w:{tag}'))
        if e is None: e=OxmlElement(f'w:{tag}'); m.append(e)
        e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa')

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; borders(t)
    hdr=t.rows[0].cells
    for i,h in enumerate(headers):
        hdr[i].text=h; shade(hdr[i],'17365D'); hdr[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in hdr[i].paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],'F3F6FA')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs: r.font.size=Pt(7.6)
    for row in t.rows:
        for i,c in enumerate(row.cells):
            set_cell_margin(c)
            if widths: c.width=Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def code(text):
    p=doc.add_paragraph(style='Code Block')
    p.paragraph_format.keep_together=True
    p.add_run(text.rstrip())
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F1F3F5'); pPr.append(shd)
    return p

def bullet(text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); return p

def step(n, title, purpose, action, result, failure=None):
    doc.add_heading(f'{n}  {title}', level=3)
    p=doc.add_paragraph(); p.add_run('Propósito  ').bold=True; p.add_run(purpose)
    p=doc.add_paragraph(); p.add_run('Acción  ').bold=True; p.add_run(action)
    p=doc.add_paragraph(); p.add_run('Resultado  ').bold=True; p.add_run(result)
    if failure:
        p=doc.add_paragraph(); p.add_run('Si falla  ').bold=True; p.add_run(failure)

def page_num(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run('Página ')
    fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text='PAGE'
    fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end')
    run._r.extend([fldChar1,instr,fldChar2])

footer=sec.footer.paragraphs[0]; footer.text='Mysic  Manual técnico de automatización  |  '; page_num(footer)
for r in footer.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,90,90)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(80)
r=p.add_run('MYSIC'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(14); r.font.color.rgb=RGBColor(23,54,93)
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Manual técnico de automatización de Mysic')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Proceso atómico  Código  Recursos  Seguridad  Recuperación'); r.font.size=Pt(12); r.font.color.rgb=RGBColor(70,70,70)
doc.add_paragraph()
table(['Elemento','Detalle'],[
    ('Proyecto','Mysic'),('Tipo','Manual técnico y procedimiento operativo'),('Ámbito','Adquisición, análisis, revisión, importación, subida y catalogación'),('Repositorio','C:\\Users\\JOSE.LAPTOP-SNQTVDID\\Desktop\\MUSICAL'),('Fecha de referencia','1 de octubre de 2026'),('Audiencia','Desarrollo, operación y administración del catálogo')],[1.6,4.9])
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(25); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Conclusión principal'); r.bold=True
p=doc.add_paragraph('La automatización está dividida intencionalmente en adquisición e importación. La primera prepara MP3; la segunda exige análisis y aprobación humana antes de etiquetar, subir o registrar contenido.', style=None); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

doc.add_heading('Propósito y alcance', level=1)
doc.add_paragraph('Este manual describe, en unidades mínimas de ejecución, cómo Mysic transforma enlaces autorizados o archivos MP3 locales en canciones reproducibles dentro del catálogo. Cada etapa identifica el disparador, el código responsable, las entradas, la operación, la salida persistida, los controles y la forma de recuperación.')
doc.add_paragraph('La implementación no automatiza la autorización legal del contenido ni la selección editorial final. La persona administradora confirma derechos, revisa metadatos, verifica la portada y ejecuta la importación definitiva.')

doc.add_heading('Mapa general de arquitectura', level=1)
code('''[React  MusicImportView]
        | invoke Tauri
        v
[Rust  src-tauri/src/lib.rs]
        | ejecuta Node
        +--> [acquisition-client.js] --> [n8n :5678] --> [worker :4310]
        |                                                | yt-dlp / FFmpeg / FFprobe
        |                                                v
        |                                  [asignar_metadatos + .ui-state.json]
        |
        +--> [ui-import.js] --> [AcoustID / MusicBrainz o Picard]
        |                              |
        +--> [run-import.js] <----------+
                 | Cloudinary          | ID3
                 | SSH + SHA-256       | catálogo
                 v                     v
          [/music en Oracle]      [PostgreSQL / API / Mysic]''')
table(['Capa','Responsabilidad','Archivo principal'],[
('React','Captura URLs y MP3, muestra estados y solicita acciones','Frontend/src/features/music-import/MusicImportView.jsx'),
('Tauri','Puente local seguro entre la UI y scripts Node','Frontend/src-tauri/src/lib.rs'),
('n8n','Orquesta creación, consulta y cancelación de trabajos','n8n/*.json'),
('Worker','Descarga, convierte, valida, deduplica y prepara MP3','automa_mysic/acquisition-worker.js'),
('Metadatos','Identifica audio y calcula destinos','automa_mysic/ui-state.js'),
('Importador','Valida, etiqueta, sube y sincroniza','automa_mysic/run-import.js'),
('Backend','Registra catálogo y asocia imágenes','Backend/scripts/*.js')],[1.0,2.35,3.1])

doc.add_heading('Recursos utilizados', level=1)
table(['Recurso','Uso','Origen o puerto'],[
('Node.js','Worker, cliente, importador y scripts del catálogo','Instalación local'),('Python 3','Dependencia requerida por yt-dlp','PATH'),('n8n','Orquestación local','127.0.0.1:5678'),('yt-dlp','Extracción de audio autorizado','Ejecutable local'),('FFmpeg','Conversión a MP3 192 kbps','Ejecutable local'),('FFprobe','Validación de duración y archivo','Ejecutable local'),('fpcalc','Huella acústica','MusicBrainz Picard'),('AcoustID','Búsqueda por huella','API HTTPS'),('MusicBrainz','Metadatos musicales','API HTTPS'),('Cloudinary','Portadas','API de administración'),('OpenSSH','Inspección y transferencia','Puerto SSH remoto'),('PostgreSQL','Catálogo de Mysic','DATABASE_URL'),('Navidrome','Identificador y reproducción','Backend')],[1.25,2.65,2.55])

doc.add_page_break()
doc.add_heading('Fase 0  Arranque de servicios', level=1)
doc.add_paragraph('Archivo responsable: scripts/start-mysic-automation.ps1')
step('0.1','Resolver rutas del proyecto','Ubicar automáticamente frontend, backend, automatización, variables y logs.','$PSScriptRoot determina la raíz y Join-Path evita concatenaciones frágiles.','Rutas absolutas internas para todos los procesos.')
code("$projectRoot = Split-Path -Parent $PSScriptRoot\n$automaRoot = Join-Path $projectRoot 'automa_mysic'\n$frontendRoot = Join-Path $projectRoot 'Frontend'\n$backendRoot = Join-Path $projectRoot 'Backend'")
step('0.2','Comprobar ejecutables','Evitar un arranque parcial con dependencias ausentes.','Require-Command y Resolve-Command buscan node, python, npm, yt-dlp, ffmpeg, ffprobe y n8n.','Ruta concreta de cada ejecutable.','El script termina con un mensaje de instalación específico.')
step('0.3','Validar secretos','Asegurar que el worker tenga autenticación real.','Lee ACQUISITION_WORKER_TOKEN y rechaza valores vacíos o de ejemplo.','Token cargado solo en memoria del proceso.')
step('0.4','Iniciar worker','Exponer el servicio de adquisición local.','Comprueba /health; si no responde, ejecuta node acquisition-worker.js oculto.','Worker listo en 127.0.0.1:4310.','El archivo logs/acquisition-worker.err.log conserva el diagnóstico.')
step('0.5','Iniciar n8n','Habilitar los webhooks de orquestación.','Comprueba /healthz; si no responde, inicia n8n y espera /healthz/readiness.','n8n listo en 127.0.0.1:5678.')
step('0.6','Iniciar backend','Habilitar autenticación y catálogo.','Ejecuta npm run dev y espera /api/health.','Backend listo en 127.0.0.1:4000.')
step('0.7','Actualizar portadas','Refrescar asociaciones antes de mostrar la aplicación.','Ejecuta music:cloudinary:apply.','Imágenes actualizadas o advertencia no bloqueante.')
step('0.8','Iniciar escritorio','Abrir Mysic y mantener unidos los procesos.','Ejecuta npm run tauri:dev, salvo -SkipDesktop.','Aplicación operativa; al cerrar, el bloque finally termina solo los procesos iniciados por el script.')

doc.add_heading('Comandos operativos', level=2)
code('''# Inicio completo
.\\scripts\\start-mysic-automation.ps1

# Solo servicios
.\\scripts\\start-mysic-automation.ps1 -SkipDesktop

# Verificación individual
Invoke-RestMethod http://127.0.0.1:5678/healthz
Invoke-RestMethod http://127.0.0.1:4000/api/health''')

doc.add_page_break()
doc.add_heading('Fase 1  Creación del trabajo de adquisición', level=1)
step('1.1','Capturar enlaces','Construir una entrada limpia desde la interfaz.','La UI divide por líneas, recorta espacios y elimina duplicados.','Arreglo de URLs únicas.')
step('1.2','Confirmar derechos','Impedir la ejecución accidental sobre contenido no autorizado.','El usuario marca rightsConfirmed; UI y worker vuelven a comprobarlo.','Valor booleano true obligatorio.')
step('1.3','Invocar Tauri','Mantener tokens fuera de React.','createMusicAcquisition invoca create_music_acquisition.','Proceso Node local con un JSON serializado.')
code('''const urls = [...new Set(sourceUrls
  .split(/\\r?\\n/)
  .map((value) => value.trim())
  .filter(Boolean))]

await musicImportApi.createMusicAcquisition(urls, true)''')
step('1.4','Autenticar contra n8n','Separar la UI del webhook local.','acquisition-client.js añade Bearer MYSIC_N8N_TOKEN y timeout de 15 segundos.','Respuesta JSON o error legible.')
step('1.5','Crear trabajo en n8n','Traducir la petición pública local a la API privada del worker.','Webhook /mysic/acquisition llama POST http://127.0.0.1:4310/jobs con el token del worker.','Trabajo aceptado y jobId devuelto.')

doc.add_heading('Contratos HTTP', level=2)
table(['Actor','Método y ruta','Entrada','Salida'],[
('Mysic a n8n','POST /webhook/mysic/acquisition','urls, rightsConfirmed','Trabajo y UUID'),
('Mysic a n8n','POST /webhook/mysic/acquisition-status','jobId','Estado y progreso'),
('Mysic a n8n','POST /webhook/mysic/acquisition-cancel','jobId','Trabajo cancelado'),
('n8n a worker','POST /jobs','urls, rightsConfirmed','HTTP 202 y trabajo'),
('n8n a worker','GET /jobs/:id','UUID en ruta','Trabajo o 404'),
('n8n a worker','POST /jobs/:id/cancel','UUID en ruta','Trabajo o 404'),
('iniciador a worker','GET /health','Bearer token','Estado de herramientas')],[1.1,2.05,1.8,1.55])

doc.add_page_break()
doc.add_heading('Fase 2  Validación y cola del worker', level=1)
step('2.1','Validar cuerpo','Limitar superficie de ataque y errores de entrada.','El servidor acepta hasta 64 KiB de JSON.','Objeto JSON interpretado.')
step('2.2','Validar URLs','Restringir cantidad, protocolo y dominio.','Máximo 20; solo http/https; host dentro de ACQUISITION_ALLOWED_HOSTS.','URLs normalizadas con new URL().')
code('''if (urls.length > 20) throw new Error('Se permiten como máximo 20 URLs por trabajo.')
if (!['http:', 'https:'].includes(parsed.protocol)) throw new Error(...)
if (!allowedHosts.has(parsed.hostname.toLowerCase())) throw new Error(...)''')
step('2.3','Persistir trabajo','Sobrevivir a reinicios.','Genera crypto.randomUUID(), estado queued y progreso 0; escribe jobs.json mediante temporal y rename.','Trabajo durable sin escritura parcial.')
step('2.4','Encolar','Serializar consumo de CPU, red y disco.','processQueue solo comienza si activeJobId está vacío.','Un trabajo activo; los demás esperan.')
step('2.5','Recuperar tras reinicio','Reanudar trabajos interrumpidos.','initialize transforma running en queued y relanza la cola.','Procesamiento recuperable.')

doc.add_heading('Modelo de estados', level=2)
table(['Estado','Significado','Transición habitual'],[
('queued','Aceptado y persistido','a running o cancelled'),('running','Procesando URLs secuencialmente','a estado terminal'),('completed','Sin errores','terminal'),('completed_with_errors','Hubo MP3 preparados y errores','terminal'),('failed','No hubo resultado utilizable','terminal'),('cancelled','Cancelación solicitada','terminal')],[1.35,3.0,2.15])

doc.add_heading('Estructura persistida', level=2)
code('''{
  "id": "UUID",
  "status": "queued",
  "progress": { "total": 2, "current": 0, "completed": 0 },
  "currentUrl": null,
  "results": [],
  "errors": [],
  "createdAt": "ISO-8601",
  "updatedAt": "ISO-8601"
}''')

doc.add_page_break()
doc.add_heading('Fase 3  Descarga y preparación del audio', level=1)
step('3.1','Crear directorio temporal','Aislar cada URL y trabajo.','Usa acquisition-data/jobs/<jobId>/<posición>.','Archivos temporales trazables.')
step('3.2','Ejecutar yt-dlp','Obtener audio autorizado en formato uniforme.','Invoca yt-dlp con configuración deshabilitada, extracción de audio y FFmpeg.','Uno o más MP3, especialmente para playlists.')
code('''yt-dlp --no-config --ignore-errors --extract-audio \\
  --audio-format mp3 --audio-quality 192K --embed-metadata \\
  --ffmpeg-location <ruta> --windows-filenames \\
  --download-archive acquisition-data/download-archive.txt \\
  --output "%(playlist_index|)02d - %(title).180B [%(id)s].%(ext)s" <URL>''')
step('3.3','Enumerar resultados','Evitar problemas de codificación de stdout en Windows.','Recorre físicamente el directorio y ordena todos los .mp3.','Lista Unicode real del sistema de archivos.')
step('3.4','Validar audio','Descartar resultados corruptos o vacíos.','FFprobe obtiene duración; fs.stat obtiene tamaño.','duration > 0 y size > 0.')
step('3.5','Calcular SHA-256','Detectar contenido idéntico aunque cambie el nombre.','Lee el archivo y calcula crypto.createHash sha256.','Hash hexadecimal.')
step('3.6','Deduplicar','Evitar dos copias de un mismo audio.','Consulta .acquisition-index.json; si existe el hash, elimina el nuevo temporal.','Resultado skipped con duplicateOf.')
step('3.7','Sanear nombre','Producir nombres válidos en Windows.','Normaliza NFC, reemplaza caracteres prohibidos, recorta y limita el stem a 180 caracteres.','Nombre seguro terminado en .mp3.')
step('3.8','Mover a bandeja','Hacer visible el archivo para la fase humana.','Busca nombre libre y renombra a asignar_metadatos/.','MP3 preparado.')
step('3.9','Registrar fila','Conectar adquisición con análisis.','addEntries añade una entrada pending en .ui-state.json.','Fila visible en Mysic.')
step('3.10','Cerrar trabajo','Consolidar resultados.','Calcula completed, completed_with_errors, failed o cancelled y persiste.','Estado terminal consultable.')

doc.add_heading('Archivos de control de adquisición', level=2)
table(['Archivo','Finalidad','Se puede reconstruir'],[
('acquisition-data/jobs.json','Estado y progreso de trabajos','Parcialmente'),('acquisition-data/download-archive.txt','IDs ya descargados por yt-dlp','Sí, pero se perdería historial'),('asignar_metadatos/.acquisition-index.json','Hash a archivo preparado','Sí, recalculando hashes'),('asignar_metadatos/.ui-state.json','Filas de revisión','Sí, con pérdida de ediciones'),('logs/mysic-errors.jsonl','Errores estructurados','No; es historial')],[2.45,2.7,1.35])

doc.add_page_break()
doc.add_heading('Fase 4  Incorporación de MP3 locales', level=1)
step('4.1','Seleccionar archivos','Permitir importación sin n8n.','Tauri abre rfd FileDialog filtrado por MP3.','Rutas locales elegidas.')
step('4.2','Preparar copias','No trabajar sobre el original del usuario.','ui-import.js copia a asignar_metadatos y omite extensiones diferentes.','Copias independientes.')
step('4.3','Resolver colisiones','No sobrescribir archivos preparados.','Añade sufijos numéricos hasta encontrar un destino libre.','Nombre único.')
step('4.4','Crear estado','Unificar el flujo manual y el adquirido.','addEntries crea exactamente la misma entrada pending.','Ambos orígenes continúan por análisis.')
code('''{
  "file": "Canción.mp3",
  "status": "pending",
  "metadata": null,
  "cloudinaryFolder": null,
  "error": null,
  "embeddedCover": false,
  "lastModified": "ISO-8601"
}''')

doc.add_heading('Fase 5  Identificación de metadatos', level=1)
step('5.1','Seleccionar candidatos','Evitar trabajo repetido.','Omite processed y ready, salvo entradas de fuente que requieren actualización.','Lista mínima de archivos a identificar.')
step('5.2','Intentar metadatos de origen','Aprovechar el identificador guardado por yt-dlp.','source-metadata extrae el ID entre corchetes y consulta la fuente.','Metadatos de fuente o activación del respaldo.')
step('5.3','Calcular huella','Identificar por contenido acústico.','fpcalc -json produce duration y fingerprint.','Huella válida o error de configuración.')
step('5.4','Consultar AcoustID','Obtener recording IDs candidatos.','Envía huella, duración y API key; aplica ACOUTID_MIN_SCORE, por defecto 0.85.','Coincidencia confiable o bloqueo.')
step('5.5','Consultar MusicBrainz','Completar estructura discográfica.','Resuelve grabación, lanzamiento, artista, año, pista, tipo y géneros.','Metadatos normalizados.')
step('5.6','Alternativa Picard','Usar selección de lanzamientos de Picard cuando está configurado.','METADATA_ENGINE=picard ejecuta SCAN y SAVE_MATCHED sobre copias temporales.','Etiquetas leídas sin modificar el original.')
step('5.7','Normalizar artista y géneros','Adaptar datos al modelo de Mysic.','Elige artista principal y limita géneros a dos, priorizando oficiales.','Campos estables.')
step('5.8','Validar japonés','Evitar nombres romanizados inventados.','Exige alias oficial entre paréntesis cuando el título contiene japonés.','Título final o estado error.')
step('5.9','Calcular destinos','Usar una convención única en Cloudinary y SSH.','deriveDestinations llama buildReleaseFolder y buildRelativeDestination.','cloudinaryFolder y remoteDestination.')
step('5.10','Persistir resultado','Mostrar el análisis y permitir corrección.','Actualiza la fila a ready o error con detalle.','Tabla de revisión actualizada.')

doc.add_page_break()
doc.add_heading('Fase 6  Revisión humana y manifiesto', level=1)
doc.add_paragraph('La interfaz recarga .ui-state.json cada cinco segundos. La persona administradora puede corregir artista, título o carpeta de Cloudinary. Guardar una corrección restablece la fila a ready.')
step('6.1','Revisar confianza','Detectar identificaciones dudosas.','La tabla muestra confidence y el error de cada fila.','Decisión informada.')
step('6.2','Confirmar lanzamiento','Verificar álbum, EP o single y su número de pista.','Se inspeccionan album, kind y trackNumber.','Ruta coherente con el catálogo.')
step('6.3','Confirmar portada','Vincular la carpeta exacta de Cloudinary.','Se conserva el valor automático o se edita manualmente.','cloudinaryFolder no vacío.')
step('6.4','Generar import.json','Dar al importador solo entradas completas.','ui-import.js import-json filtra status ready y metadatos obligatorios.','Manifiesto mínimo y reproducible.')
code('''{
  "imports": [{
    "file": "01 - Canción.mp3",
    "metadataSource": "musicbrainz",
    "artist": "Artista",
    "title": "Canción",
    "kind": "album",
    "releaseName": "Álbum",
    "year": 2026,
    "genres": ["Pop", "Rock"],
    "trackNumber": 1,
    "musicbrainzRecordingId": "...",
    "cloudinaryFolder": "Artistas/Artista/Albums/Álbum"
  }]
}''')

doc.add_heading('Reglas de rutas', level=2)
table(['Tipo','Carpeta','Archivo final'],[
('album','Artistas/<artista>/Albums/<lanzamiento>','NN - <título>.mp3'),('ep','Artistas/<artista>/EPs/<lanzamiento>','NN - <título>.mp3'),('single','Artistas/<artista>/Singles/<lanzamiento>','NN - <título>.mp3')],[1.0,3.65,1.85])

doc.add_page_break()
doc.add_heading('Fase 7  Vista previa no destructiva', level=1)
code('node run-import.js --dry-run --remote-preview')
step('7.1','Leer manifiesto','Revalidar fuera de la UI.','readManifest analiza JSON y normaliza cada entrada.','Entradas tipadas y saneadas.')
step('7.2','Validar segmentos','Bloquear traversal y separadores.','Rechaza vacío, punto, doble punto, slash y backslash en segmentos.','Ruta confinada.')
step('7.3','Actualizar metadatos','Resolver nuevamente cuando metadataSource es musicbrainz.','resolveMusicBrainzEntry completa la entrada efectiva.','Metadatos finales.')
step('7.4','Listar Cloudinary','Obtener activos disponibles.','Usa CLOUDINARY_CLOUD_NAME, API_KEY, API_SECRET y modo dynamic o fixed.','Lista de imágenes.')
step('7.5','Elegir portada','Exigir coincidencia exacta de carpeta.','Filtra formatos válidos y HTTPS; prioriza cover, portada y luego nombre.','Una portada o fallo de la fila.')
step('7.6','Inspeccionar SSH','Detectar conflicto antes de escribir.','Comprueba carpeta y archivo mediante ssh en modo BatchMode.','ready si no existe; conflict si existe.')
step('7.7','Devolver resumen','Permitir decisión humana.','No etiqueta, mueve ni sube ningún archivo.','JSON con artista, título, portada, carpeta y destino.')

doc.add_heading('Selección de portada', level=2)
code('''const matches = assets
  .filter((asset) => getAssetFolder(asset, folderMode) === folder)
  .filter((asset) => IMAGE_FORMATS.has(asset.format) && asset.secure_url.startsWith('https://'))
  .sort(preferirCoverOPortadaLuegoNombre)

if (!matches.length) throw new Error(`La carpeta ${folder} no contiene una imagen válida.`)''')

doc.add_page_break()
doc.add_heading('Fase 8  Importación y etiquetado', level=1)
code('node run-import.js --sync-catalog --upload-remote')
step('8.1','Comprobar idempotencia','No repetir entradas sin cambios.','Calcula SHA-256 del JSON normalizado y compara .state.json.','skipped si el hash coincide.')
step('8.2','Calcular destino local','Construir el mismo árbol que usa el servidor.','buildDestination une MUSIC_LIBRARY_PATH y destinationParts.','Ruta local absoluta confinada.')
step('8.3','Comprobar conflictos','No sobrescribir una canción.','fs.stat rechaza si el destino ya existe.','Error antes de modificar el original.')
step('8.4','Crear copia temporal','Proteger el MP3 de entrada.','Copia a <destino>.tmp-<pid>-<fecha>.','Archivo de trabajo aislado.')
step('8.5','Detectar portada embebida','Conservar arte existente.','music-metadata inspecciona common.picture.','Decisión de descargar o no portada.')
step('8.6','Descargar portada','Añadir arte solo cuando falta.','fetch HTTPS con timeout de 20 segundos.','Buffer de imagen.')
step('8.7','Escribir ID3','Incorporar metadatos completos.','node-id3 actualiza título, artista, álbum, año, género, pista, tipo y portada.','Temporal etiquetado.')
step('8.8','Publicar localmente','Hacer atómica la colocación local.','Renombra el temporal al destino.','MP3 final local; el original aún se conserva si habrá subida remota.')

doc.add_heading('Etiquetas ID3 escritas', level=2)
table(['Etiqueta','Fuente'],[
('title','entry.title'),('artist','entry.artist'),('album','entry.releaseName'),('year','entry.year'),('genre','hasta dos géneros unidos por punto y coma'),('trackNumber','entry.trackNumber'),('MusicBrainz Album Type','entry.kind'),('front cover','Cloudinary solo si no había imagen')],[2.15,4.35])

doc.add_page_break()
doc.add_heading('Fase 9  Subida SSH verificada', level=1)
step('9.1','Validar configuración','Evitar comandos con host, usuario o raíz inseguros.','getSshConfig valida MUSIC_SSH_HOST, USER, KEY_PATH y ROOT.','Configuración SSH normalizada.')
step('9.2','Calcular hash local','Crear la expectativa de integridad.','Streaming SHA-256 evita cargar el archivo completo.','expectedHash.')
step('9.3','Comprobar destino','Prohibir sobrescritura remota.','El shell remoto sale con código 73 si el destino existe.','Importación abortada sin cambios.')
step('9.4','Crear carpeta','Preparar la jerarquía remota.','mkdir -p -- <directorio>.','Carpeta disponible.')
step('9.5','Transmitir temporal','Evitar archivos definitivos parciales.','El flujo del MP3 entra por stdin a <destino>.upload-*.part.','Archivo temporal remoto.')
step('9.6','Verificar integridad','Detectar corrupción de transferencia.','sha256sum remoto debe coincidir con expectedHash.','Transferencia validada.')
step('9.7','Renombrar atómicamente','Publicar solo un archivo completo.','mv cambia el .part al destino definitivo.','MP3 visible en /music.')
step('9.8','Limpiar origen','Cerrar la operación local.','Solo después de subir, elimina el MP3 de asignar_metadatos.','Sin copia pendiente innecesaria.')
code('''set -eu
if [ -e "$destination" ]; then exit 73; fi
mkdir -p -- "$directory"
trap 'rm -f -- "$temporary"' 0 1 2 15
cat > "$temporary"
actual_hash=$(sha256sum -- "$temporary" | cut -d ' ' -f 1)
[ "$actual_hash" = "$expected_hash" ]
mv -- "$temporary" "$destination"
trap - 0 1 2 15''')

doc.add_heading('Semántica ante fallos', level=2)
table(['Fallo','Efecto','Recuperación'],[
('Destino remoto existente','No se sobrescribe','Cambiar metadatos/ruta o retirar conflicto'),('SSH interrumpido','Se elimina .part por trap','Reintentar'),('Hash distinto','No se ejecuta mv','Reintentar e investigar red/disco'),('Error tras etiquetar','Se elimina destino local generado','El original preparado se conserva'),('Portada inaccesible','No se publica la canción','Corregir Cloudinary y reintentar')],[1.7,2.35,2.45])

doc.add_page_break()
doc.add_heading('Fase 10  Sincronización del catálogo', level=1)
step('10.1','Registrar resultado','Marcar idempotencia.','Guarda hash de la entrada en .state.json después de procesar y subir.','La misma entrada se omitirá en ejecuciones futuras.')
step('10.2','Evaluar tanda','Evitar catalogar una operación global incompleta.','Solo sincroniza si failed es cero y hay processed o skipped.','Tanda apta para catálogo.')
step('10.3','Importar estructura','Crear o actualizar entidades musicales.','Ejecuta Backend/scripts/import-music.js <libraryRoot> --apply.','PostgreSQL refleja artista, lanzamiento y pista.')
step('10.4','Sincronizar imágenes','Asociar recursos visuales.','Ejecuta Backend/scripts/sync-cloudinary-images.js --apply.','Portadas y perfiles enlazados.')
step('10.5','Refrescar UI','Mostrar resultado sin reiniciar.','React ejecuta onCatalogRefresh y vuelve a leer el estado.','Canción disponible en Mysic.')

doc.add_heading('Condición exacta de sincronización', level=2)
code('''if (!dryRun && syncCatalog && summary.failed === 0
    && summary.processed + summary.skipped > 0) {
  await runSync('music:import', libraryRoot)
  await runSync('music:cloudinary:apply', libraryRoot)
}''')
doc.add_paragraph('Implicación operativa: si una fila falla, otras pueden haber sido subidas, pero la sincronización global queda pendiente. Se debe corregir la fila, reintentar y confirmar el resultado del catálogo.')

doc.add_heading('Variables de configuración', level=1)
table(['Archivo','Variable','Función'],[
('automa_mysic/.env','ACQUISITION_WORKER_TOKEN','Autentica n8n frente al worker'),('automa_mysic/.env','MYSIC_N8N_TOKEN','Autentica Mysic frente a n8n'),('automa_mysic/.env','ACQUISITION_WORKER_PORT','Puerto del worker, 4310 por defecto'),('automa_mysic/.env','ACQUISITION_ALLOWED_HOSTS','Lista de dominios permitidos'),('automa_mysic/.env','MYSIC_N8N_WEBHOOK_BASE','Base de webhooks'),('Backend/.env','FPCALC_PATH','Ejecutable de huella'),('Backend/.env','PICARD_PATH','Ejecutable Picard opcional'),('Backend/.env','METADATA_ENGINE','musicbrainz o picard'),('Backend/.env','ACOUSTID_API_KEY','Credencial de identificación'),('Backend/.env','ACOUSTID_MIN_SCORE','Umbral de confianza'),('Backend/.env','MUSIC_LIBRARY_PATH','Biblioteca local temporal'),('Backend/.env','MUSIC_SSH_HOST USER KEY_PATH ROOT','Destino remoto'),('Backend/.env.cloudinary','CLOUDINARY variables','Consulta de imágenes')],[1.55,2.25,2.7])

doc.add_page_break()
doc.add_heading('Seguridad y controles', level=1)
table(['Control','Implementación','Riesgo mitigado'],[
('Doble consentimiento','UI y worker exigen rightsConfirmed','Uso accidental no autorizado'),('Bind local','Worker escucha solo 127.0.0.1','Exposición a la red'),('Dos tokens','Mysic→n8n y n8n→worker','Movimiento lateral'),('Allowlist de hosts','Dominios configurables','SSRF y fuentes no previstas'),('Límite de 20 URLs','validateUrls','Abuso de recursos'),('Cuerpo de 64 KiB','readBody','Memoria y payload excesivo'),('Rutas saneadas','cleanSegment, normalizeFolder, isInside','Traversal y escritura fuera de raíz'),('No sobrescritura','Comprobaciones local y remota','Pérdida de música existente'),('Temporales y rename','JSON, MP3 y subida','Archivos parciales'),('SHA-256','Deduplicación e integridad SSH','Duplicados y corrupción'),('BatchMode SSH','Sin prompt interactivo','Bloqueos invisibles'),('Redacción de secretos','error-log reemplaza campos sensibles','Filtración en logs')],[1.55,3.1,1.85])

doc.add_heading('Observabilidad y diagnóstico', level=1)
table(['Recurso','Contenido','Cuándo revisarlo'],[
('logs/acquisition-worker.err.log','Inicio y fallos del worker','/health no responde'),('logs/n8n.err.log','Problemas de n8n','Webhook no disponible'),('logs/backend.err.log','Backend y catálogo','/api/health o catálogo falla'),('logs/mysic-desktop.err.log','Tauri y Vite','La aplicación no abre'),('logs/mysic-errors.jsonl','Errores estructurados por componente','Fallos por canción o trabajo'),('acquisition-data/jobs/<id>/yt-dlp.log','stdout y stderr de yt-dlp','Adquisición parcial o fallida')],[2.35,2.4,1.75])

doc.add_heading('Procedimiento de diagnóstico', level=2)
for i,t in enumerate([
'Confirmar que el iniciador sigue activo y revisar los cuatro logs de proceso.',
'Consultar /health del worker usando el Bearer token y verificar yt-dlp, FFmpeg, FFprobe e inbox.',
'Consultar el trabajo por UUID y comparar status, currentUrl, progress, results y errors.',
'Buscar el evento correspondiente en mysic-errors.jsonl por timestamp, source o jobId.',
'Comprobar .ui-state.json para saber si el MP3 llegó a la fase de revisión.',
'Ejecutar Vista previa para aislar errores de metadatos, Cloudinary o destino SSH.',
'Reintentar solo después de corregir la causa; los controles de idempotencia evitan duplicados.'
],1):
    doc.add_paragraph(f'{i}. {t}')

doc.add_page_break()
doc.add_heading('Recuperación por escenario', level=1)
table(['Escenario','Evidencia','Acción segura'],[
('Worker reiniciado durante descarga','Trabajo vuelve a queued','Dejar que processQueue lo retome'),('MP3 quedó en directorio temporal','Existe bajo acquisition-data/jobs/<id>','Ejecutar recover-acquisition.js <jobId>'),('Fila en error de análisis','.ui-state.json contiene error','Corregir configuración o editar metadatos y reanalizar'),('Carpeta Cloudinary incorrecta','Vista previa indica que no hay imagen','Subir o corregir carpeta y repetir preview'),('Conflicto remoto','remote.fileExists true','No forzar; verificar si la canción ya existe'),('Subida cortada','No hay destino final; .part se limpia','Reintentar importación'),('Archivo subido pero catálogo pendiente','Tanda tuvo al menos un failed','Corregir y reintentar con sync-catalog'),('Entrada marcada como procesada','Hash coincide en .state.json','Modificar la entrada solo si realmente debe reprocesarse')],[1.55,2.35,2.85])

doc.add_heading('Comandos de mantenimiento', level=1)
code('''# Pruebas del importador
cd automa_mysic
npm test

# Vista previa por CLI
npm run preview

# Importación local y catálogo
npm run import -- --sync-catalog

# Recuperación de archivos de un trabajo
node recover-acquisition.js <jobId>

# Catálogo remoto y portadas
cd ..\\Backend
npm run music:filezilla
npm run music:filezilla:apply
npm run music:cloudinary
npm run music:cloudinary:apply''')

doc.add_heading('Criterios de finalización', level=1)
for t in [
'El trabajo de adquisición está en un estado terminal y no tiene archivos temporales necesarios sin recuperar.',
'Cada canción aprobada tiene metadatos completos, tipo de lanzamiento y número de pista.',
'La vista previa no informa conflicto remoto ni portada ausente.',
'La subida devuelve destino, tamaño y SHA-256.',
'La entrada queda registrada en .state.json.',
'Las sincronizaciones de catálogo y Cloudinary terminan sin error.',
'La canción aparece en /api/catalog y puede reproducirse mediante el identificador de Navidrome.'
]: bullet(t)

doc.add_heading('Inventario de código', level=1)
table(['Archivo','Responsabilidad'],[
('scripts/start-mysic-automation.ps1','Arranque, salud, logs y cierre'),('Frontend/src/features/music-import/MusicImportView.jsx','Interacción y polling'),('Frontend/src/features/music-import/musicImportApi.js','Invocaciones Tauri'),('Frontend/src-tauri/src/lib.rs','Puente Rust a Node'),('automa_mysic/acquisition-client.js','Cliente autenticado de n8n'),('n8n/mysic-acquisition-workflow.json','Crear y observar trabajos'),('n8n/mysic-acquisition-status-workflow.json','Consultar progreso'),('n8n/mysic-acquisition-cancel-workflow.json','Cancelar'),('automa_mysic/acquisition-worker.js','Adquisición y staging'),('automa_mysic/ui-import.js','Preparación, análisis y manifiesto'),('automa_mysic/ui-state.js','Estado e identificación'),('automa_mysic/musicbrainz.js','AcoustID y MusicBrainz'),('automa_mysic/picard.js','Motor Picard opcional'),('automa_mysic/import-music.js','Validación, rutas y portada'),('automa_mysic/process-entry.js','ID3 y publicación local'),('automa_mysic/remote-library.js','SSH, hash y publicación remota'),('automa_mysic/run-import.js','Coordinación final'),('automa_mysic/error-log.js','Registro y redacción'),('Backend/scripts/import-music.js','Persistencia del catálogo'),('Backend/scripts/sync-cloudinary-images.js','Asociación de imágenes')],[3.55,3.0])

doc.add_heading('Resumen de garantías', level=1)
doc.add_paragraph('El diseño garantiza que adquirir no significa publicar. Un archivo atraviesa una bandeja de preparación, identificación automática, revisión humana y vista previa no destructiva antes de que se escriban etiquetas o se cree un destino remoto. La publicación usa temporales, no sobrescribe archivos, verifica SHA-256 y conserva el original cuando falla. Finalmente, el catálogo solo se sincroniza bajo una condición explícita y observable.')

# Metadata and save
doc.core_properties.title='Manual técnico de automatización de Mysic'
doc.core_properties.subject='Proceso atómico, código, recursos, seguridad y recuperación'
doc.core_properties.author='Proyecto Mysic'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
